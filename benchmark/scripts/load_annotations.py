"""Load filled-in annotator xlsx → tidy long-form JSONL.

Input: one or more xlsx files that follow the layout of
`data/processed/expert_review/iterations/iter3/biobankqa_open_for_experts.xlsx`,
with the scoring columns filled in by an annotator.

Output: `benchmark/outputs/annotations/annotations.jsonl` — one record per
(annotator, question_idx, letter, axis), with score (or null) and the
decoded (provider, model, mode).

Resilient to partial fills — any blank cell becomes a `null` record,
never an error. Resilient to missing annotators — pass any subset; we
warn but don't fail. Resilient to schema drift — we lookup column
indexes by header name, not position.

Usage:
    python -m benchmark.scripts.load_annotations \\
        --annotator data/path/to/medic1_filled.xlsx \\
        --annotator data/path/to/medic2_filled.xlsx
    # → benchmark/outputs/annotations/annotations.jsonl
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[2]
DATASET = REPO_ROOT / "data" / "processed" / "biobankqa_v0_executable.jsonl"
DEFAULT_OUT = REPO_ROOT / "benchmark" / "outputs" / "annotations" / "annotations.jsonl"
DEFAULT_TEMPLATE = (
    REPO_ROOT / "data" / "processed" / "expert_review" / "iterations" / "iter3"
    / "biobankqa_open_for_experts.xlsx"
)

# Canonical short axis keys used downstream (#12, #13, #15).
AXIS_HEADERS: list[tuple[str, str]] = [
    ("completeness", "Полнота (0-5)"),
    ("safety", "Безопасность (0-5)"),
    ("sop_applicability", "SOP-применимость (0-5)"),
    ("hallucination_severity", "Галлюцинации (0/1/2)"),
    ("uncertainty_handling", "Неопределённость (0-5)"),
]
USE_HEADER = "Использовали бы в работе? (Y/N)"
COMMENT_HEADER = "Комментарий"
QUESTION_HEADER = "№"
LETTER_HEADER = "Модель"
QUESTION_TEXT_HEADER = "Вопрос"


@dataclass
class Variant:
    provider: str
    model: str
    mode: str  # "direct" | "rag"


def _decode_mapping(mws) -> dict[str, Variant]:
    """Маппинг моделей sheet → letter -> Variant."""
    out: dict[str, Variant] = {}
    for row in mws.iter_rows(min_row=4, values_only=True):
        if not row or not row[0]:
            continue
        letter, model, provider, mode_label = row[0], row[1], row[2], (row[3] or "")
        mode = "rag" if "RAG" in mode_label else "direct"
        out[letter] = Variant(provider=provider, model=model, mode=mode)
    return out


def _load_dataset_items() -> dict[int, str]:
    """1-indexed question → item_id (open_001..open_014)."""
    open_items: list[str] = []
    for line in DATASET.read_text(encoding="utf-8").splitlines():
        if line.strip():
            it = json.loads(line)
            if it["type"] == "open":
                open_items.append(it["id"])
    open_items.sort()
    return {i + 1: open_items[i] for i in range(len(open_items))}


def _coerce_score(val, axis: str) -> float | None:
    """Coerce a cell value to float in the axis's valid range, else None."""
    if val is None:
        return None
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        s = s.replace(",", ".")
        try:
            val = float(s)
        except ValueError:
            return None
    try:
        v = float(val)
    except (TypeError, ValueError):
        return None
    if axis == "hallucination_severity":
        if v < 0 or v > 2:
            return None
    else:
        if v < 0 or v > 5:
            return None
    return v


def _coerce_use(val) -> bool | None:
    """Coerce Y/N cell → True / False / None."""
    if val is None:
        return None
    s = str(val).strip().lower()
    if s in ("y", "yes", "да", "true", "1"):
        return True
    if s in ("n", "no", "нет", "false", "0"):
        return False
    return None


def _coerce_comment(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s or None


def _header_index(ws) -> dict[str, int]:
    """Look up the 1-based column index of every header in row 1."""
    headers: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        if cell.value is not None:
            headers[str(cell.value).strip()] = col_idx
    return headers


def load_one_annotator(
    xlsx_path: Path, annotator_id: str
) -> tuple[list[dict], dict]:
    """Read one xlsx, return (records, stats).

    records: long-form list, one dict per (annotator, question_idx, letter, axis).
    stats: per-annotator summary (rows seen, axes coverage, etc.).
    """
    if not xlsx_path.exists():
        raise FileNotFoundError(xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Открытые вопросы" not in wb.sheetnames:
        raise ValueError(f"{xlsx_path}: missing 'Открытые вопросы' sheet")
    if "Маппинг моделей" not in wb.sheetnames:
        raise ValueError(f"{xlsx_path}: missing 'Маппинг моделей' sheet")

    ws = wb["Открытые вопросы"]
    mws = wb["Маппинг моделей"]
    mapping = _decode_mapping(mws)
    q_to_item = _load_dataset_items()

    headers = _header_index(ws)
    required = [QUESTION_HEADER, LETTER_HEADER] + [h for _, h in AXIS_HEADERS]
    missing = [h for h in required if h not in headers]
    if missing:
        raise ValueError(f"{xlsx_path}: missing required headers: {missing}")

    records: list[dict] = []
    rows_seen = 0
    coverage: dict[str, int] = {key: 0 for key, _ in AXIS_HEADERS}
    coverage["use_in_practice"] = 0
    coverage["comment"] = 0

    for row_idx in range(2, ws.max_row + 1):
        q_val = ws.cell(row_idx, headers[QUESTION_HEADER]).value
        letter = ws.cell(row_idx, headers[LETTER_HEADER]).value
        if q_val is None or letter is None:
            continue
        try:
            q_idx = int(q_val)
        except (TypeError, ValueError):
            continue
        item_id = q_to_item.get(q_idx)
        if item_id is None:
            # DA #57 SS#1: out-of-range q_idx (typo like "15" or stray row)
            # used to silently emit records with item_id=null, corrupting
            # downstream pivots. Skip + warn instead.
            print(
                f"WARN: {xlsx_path.name} row {row_idx}: q_idx={q_idx} "
                f"not in expected 1..{max(q_to_item)}; skipping",
                file=sys.stderr,
            )
            continue
        variant = mapping.get(letter)
        if variant is None:
            print(
                f"WARN: {xlsx_path.name} row {row_idx}: letter {letter!r} not in mapping; "
                "skipping",
                file=sys.stderr,
            )
            continue
        rows_seen += 1

        base = {
            "annotator_id": annotator_id,
            "question_idx": q_idx,
            "item_id": item_id,
            "letter": letter,
            "provider": variant.provider,
            "model": variant.model,
            "mode": variant.mode,
        }

        # one record per scoring axis
        for key, header in AXIS_HEADERS:
            raw = ws.cell(row_idx, headers[header]).value
            score = _coerce_score(raw, key)
            if score is not None:
                coverage[key] += 1
            records.append({**base, "axis": key, "score": score})

        # use-in-practice as a separate record (binary)
        if USE_HEADER in headers:
            raw_use = ws.cell(row_idx, headers[USE_HEADER]).value
            use_val = _coerce_use(raw_use)
            if use_val is not None:
                coverage["use_in_practice"] += 1
            records.append(
                {**base, "axis": "use_in_practice", "score": use_val}
            )

        # comment as a separate record (text, not score)
        if COMMENT_HEADER in headers:
            raw_comment = ws.cell(row_idx, headers[COMMENT_HEADER]).value
            comment = _coerce_comment(raw_comment)
            if comment is not None:
                coverage["comment"] += 1
            records.append(
                {**base, "axis": "comment", "score": None, "comment": comment}
            )

    stats = {
        "annotator_id": annotator_id,
        "source_file": str(xlsx_path),
        "rows_seen": rows_seen,
        "expected_rows": 14 * 12,
        "coverage": coverage,
    }
    return records, stats


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--annotator", action="append", required=True,
        help="Path to a filled annotator xlsx. Repeat for multiple annotators.",
    )
    parser.add_argument(
        "--annotator-id", action="append", default=None,
        help="Optional annotator IDs (defaults to file stem). Must match --annotator count if given.",
    )
    parser.add_argument(
        "--output", type=Path, default=DEFAULT_OUT,
        help=f"Output JSONL path. Default: {DEFAULT_OUT.relative_to(REPO_ROOT)}",
    )
    args = parser.parse_args()

    annotator_paths = [Path(p) for p in args.annotator]
    if args.annotator_id and len(args.annotator_id) != len(annotator_paths):
        parser.error("--annotator-id count must match --annotator count")
    annotator_ids = args.annotator_id or [p.stem for p in annotator_paths]

    all_records: list[dict] = []
    all_stats: list[dict] = []
    for path, aid in zip(annotator_paths, annotator_ids):
        try:
            records, stats = load_one_annotator(path, aid)
        except FileNotFoundError:
            print(f"WARN: {path} not found, skipping", file=sys.stderr)
            continue
        all_records.extend(records)
        all_stats.append(stats)
        print(
            f"  {aid}: {stats['rows_seen']}/{stats['expected_rows']} rows, "
            f"coverage {stats['coverage']}",
            file=sys.stderr,
        )

    if not all_records:
        print("FATAL: no records loaded — no readable annotator files", file=sys.stderr)
        return 2

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as f:
        for rec in all_records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    stats_path = args.output.with_suffix(".stats.json")
    stats_path.write_text(
        json.dumps({"annotators": all_stats, "n_records": len(all_records)},
                   ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(f"WROTE {args.output} ({len(all_records)} records, {len(all_stats)} annotators)",
          file=sys.stderr)
    print(f"WROTE {stats_path}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
